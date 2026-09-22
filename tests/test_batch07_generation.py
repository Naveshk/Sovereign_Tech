
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from app.services.file_generator import generate_approval_note_docx
from app.services.file_tools.excel_tools import generate_excel_from_records
from app.services.file_tools.validation_tools import validate_generated_file

def test_approval_note_has_review_controls(tmp_path, monkeypatch):
    import app.services.file_generator as fg
    monkeypatch.setattr(fg, "GENERATED_DIR", str(tmp_path))
    out = generate_approval_note_docx(
        "## Scope\n- Review the maintenance request.",
        title="MRPL Approval Note",
        requester="EMP001",
        evidence_refs=[{"source":"manual.pdf","page":"14","chunk_index":"2","classification":"internal"}],
    )
    doc = Document(out["file_path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Pending human approval" in text
    assert "Decision:" in text
    assert "Reviewer / Approver:" in text
    assert any(t.cell(0,0).text == "Source" for t in doc.tables)

def test_xlsx_generation_includes_summary_and_exact_rows(tmp_path):
    path = tmp_path / "report.xlsx"
    out = generate_excel_from_records(
        str(path),
        [{"Tag":"P-101","Pressure":10.5},{"Tag":"P-102","Pressure":12.0}],
        title="Inspection Data",
        summary={"Source":"inspection.csv","Extraction method":"source_workbook_rows","Records":2},
    )
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["Data", "Summary"]
    assert wb["Data"]["A3"].value == "P-101"
    assert wb["Data"]["B3"].value == 10.5
    assert wb["Summary"]["A3"].value == "Source"
    wb.close()
    validate_generated_file(str(path), "xlsx")

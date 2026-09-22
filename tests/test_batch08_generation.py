
from pathlib import Path
from pptx import Presentation
import fitz

from app.services.file_tools.ppt_tools import generate_pptx
from app.services.file_tools.document_generators import markdown_to_pdf
from app.services.file_tools.validation_tools import validate_generated_file


def test_pptx_generation_renders_real_table_and_bullets(tmp_path):
    path = tmp_path / "report.pptx"
    content = """# Findings
- Pressure is within the supplied limit.
- Review required before approval.

# Measurements
| Tag | Pressure | Status |
|---|---:|---|
| P-101 | 10.5 | Review |
| P-102 | 12.0 | OK |
"""
    out = generate_pptx(str(path), "Inspection Report", content, "inspection.pdf")
    prs = Presentation(path)
    assert out["slides"] == 3
    assert prs.slides[0].shapes.title.text == "Inspection Report"
    assert any(getattr(shape, "has_table", False) for shape in prs.slides[2].shapes)
    validate_generated_file(str(path), "pptx")


def test_pdf_generation_creates_openable_pages_and_title(tmp_path, monkeypatch):
    import app.services.file_tools.document_generators as dg
    monkeypatch.setattr(dg, "GENERATED_DIR", tmp_path)
    path_info = markdown_to_pdf(
        "# Findings\nThe inspection result is reviewable.\n\n## Data\n| Tag | Value |\n|---|---:|\n| P-101 | 10.5 |",
        "Inspection Report",
    )
    doc = fitz.open(path_info["file_path"])
    assert len(doc) >= 1
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    assert "Inspection Report" in text
    validate_generated_file(path_info["file_path"], "pdf")

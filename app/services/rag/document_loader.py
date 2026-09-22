from pathlib import Path
import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from app.services.ocr_service import local_ocr


def load_document(path: str) -> list[dict]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == '.pdf':
        return _load_pdf(p)
    if ext == '.docx':
        return _load_docx(p)
    if ext in {'.txt', '.md', '.csv'}:
        return [{'page': 1, 'text': p.read_text(encoding='utf-8', errors='ignore')}]
    if ext == '.xlsx':
        return _load_xlsx(p)
    if ext == '.pptx':
        return _load_pptx(p)
    raise ValueError(f'RAG ingestion does not support {ext}. Use PDF, DOCX, PPTX, XLSX, TXT, MD or CSV.')


def _load_pdf(path: Path) -> list[dict]:
    pages = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc):
            text = page.get_text('text').strip()
            if not text:
                # Keep RAG useful for scanned PDFs by reusing the existing local OCR path.
                pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False)
                with __import__("tempfile").TemporaryDirectory(prefix="sovereign_rag_pdf_") as tmp:
                    image_path = Path(tmp) / f"page_{i + 1}.png"
                    pix.save(str(image_path))
                    ocr = local_ocr.extract(image_path)
                    text = (ocr.get("text") or "").strip()
            pages.append({'page': i + 1, 'text': text})
    return pages


def _load_docx(path: Path) -> list[dict]:
    doc = Document(path)
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(' | '.join(cell.text.strip() for cell in row.cells))
    return [{'page': 1, 'text': '\n'.join(lines)}]

def _load_xlsx(path: Path) -> list[dict]:
    pages = []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for index, ws in enumerate(wb.worksheets, 1):
            rows = []
            for row in ws.iter_rows(values_only=True):
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values:
                    rows.append(' | '.join(values))
            if rows:
                pages.append({'page': index, 'text': f'Sheet: {ws.title}\n' + '\n'.join(rows)})
    finally:
        wb.close()
    return pages


def _load_pptx(path: Path) -> list[dict]:
    prs = Presentation(path)
    pages = []
    for index, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                lines.append(shape.text.strip())
            if getattr(shape, 'has_table', False):
                for row in shape.table.rows:
                    lines.append(' | '.join(cell.text.strip() for cell in row.cells))
        if lines:
            pages.append({'page': index, 'text': '\n'.join(lines)})
    return pages

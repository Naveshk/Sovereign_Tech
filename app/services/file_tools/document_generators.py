import os, re
from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment


BASE_DIR = Path(__file__).resolve().parents[3]
GENERATED_DIR = BASE_DIR / "data" / "generated"
GENERATED_DIR.mkdir(parents=True, exist_ok=True)


def safe_stem(title: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "_", title or "generated_file").strip("_").lower()
    return s or "generated_file"


def _split_markdown_tables(content: str) -> list[list[list[str]]]:
    blocks, current = [], []
    for line in (content or "").splitlines():
        line = line.strip()
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if cells and not all(set(c) <= set("-: ") for c in cells):
                current.append(cells)
        else:
            if current:
                blocks.append(current)
                current = []
    if current:
        blocks.append(current)
    return blocks


def _pdf_page_number(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(A4[0] - 42, 22, f"Page {doc.page}")
    canvas.restoreState()


def markdown_to_pdf(content: str, title: str) -> dict:
    filename = safe_stem(title) + ".pdf"
    path = GENERATED_DIR / filename
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CenterTitle", parent=styles["Title"], alignment=TA_CENTER, spaceAfter=16))
    story = [Paragraph(title, styles["CenterTitle"])]

    # Render text and Markdown tables as real PDF tables.
    lines = (content or "").splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i].strip()
        if not raw:
            story.append(Spacer(1, 6)); i += 1; continue

        if raw.startswith("|") and raw.endswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    block.append(cells)
                i += 1
            if len(block) >= 2:
                table = Table(block, repeatRows=1)
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#334155")),
                    ("TEXTCOLOR", (0,0), (-1,0), colors.white),
                    ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
                    ("GRID", (0,0), (-1,-1), 0.4, colors.grey),
                    ("VALIGN", (0,0), (-1,-1), "TOP"),
                    ("FONTSIZE", (0,0), (-1,-1), 8),
                    ("LEFTPADDING", (0,0), (-1,-1), 5),
                    ("RIGHTPADDING", (0,0), (-1,-1), 5),
                ]))
                story.extend([table, Spacer(1, 10)])
            continue

        safe = raw.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if raw.startswith("### "): story.append(Paragraph(safe[4:], styles["Heading3"]))
        elif raw.startswith("## "): story.append(Paragraph(safe[3:], styles["Heading2"]))
        elif raw.startswith("# "): story.append(Paragraph(safe[2:], styles["Heading1"]))
        elif raw.startswith("- "): story.append(Paragraph("• " + safe[2:], styles["BodyText"]))
        else: story.append(Paragraph(safe, styles["BodyText"]))
        story.append(Spacer(1, 4))
        i += 1

    SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=42, leftMargin=42,
        topMargin=42, bottomMargin=36,
        title=title, author="Sovereign AI Workbench"
    ).build(story, onFirstPage=_pdf_page_number, onLaterPages=_pdf_page_number)

    return {"filename": filename, "file_path": str(path), "file_type": "pdf", "title": title}


# Legacy helper retained for callers that used this module directly.
def markdown_tables(content: str) -> list[list]:
    blocks = _split_markdown_tables(content)
    return blocks[0] if blocks else []


def generate_excel(
    title: str,
    analysis: dict | None = None,
    rows: list[list] | None = None,
    content: str | None = None
) -> dict:
    """Legacy Excel interface retained for backward compatibility.

    New agent code uses excel_tools.generate_excel_from_records so structured
    records, rather than free-form LLM prose, drive the workbook.
    """
    from app.services.file_tools.excel_tools import generate_excel_from_records, markdown_table_to_records

    filename = safe_stem(title) + ".xlsx"
    path = GENERATED_DIR / filename

    if rows:
        headers = rows[0]
        records = [dict(zip(headers, row)) for row in rows[1:] if any(str(x).strip() for x in row)]
    elif content:
        records = markdown_table_to_records(content)
    else:
        records = []

    if records:
        return generate_excel_from_records(str(path), records, sheet_name="Data", title=title)

    # Preserve legacy behavior when only workbook analysis is supplied.
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16)
    if analysis:
        ws["A3"], ws["B3"] = "Workbook", analysis.get("file", "")
        ws["A4"], ws["B4"] = "Sheets", analysis.get("sheet_count", 0)
        row = 6
        for sh in analysis.get("sheets", []):
            ws.cell(row, 1, sh["name"])
            ws.cell(row, 2, f"{sh['rows']} rows × {len(sh['columns'])} columns")
            row += 1
    elif content:
        data = wb.create_sheet("AI Response")
        for line in content.splitlines():
            if line.strip():
                data.append([line.strip()])
    wb.save(path)
    check = load_workbook(path, read_only=True)
    try:
        if not check.sheetnames:
            raise ValueError("Generated workbook is invalid.")
    finally:
        check.close()
    return {"filename": filename, "file_path": str(path), "file_type": "xlsx", "title": title}

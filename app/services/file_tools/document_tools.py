from __future__ import annotations

from pathlib import Path
from docx import Document


def analyze_docx(path: str) -> dict:
    p = Path(path)
    doc = Document(p)
    paragraphs = []
    headings = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        item = {"text": text, "style": para.style.name if para.style else ""}
        paragraphs.append(item)
        if item["style"].lower().startswith("heading"):
            headings.append(text)

    tables = []
    for index, table in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append({"index": index, "rows": rows})

    return {
        "file": p.name,
        "paragraphs": paragraphs,
        "headings": headings,
        "tables": tables,
        "text": "\n".join(x["text"] for x in paragraphs),
    }

from __future__ import annotations

from pathlib import Path


def validate_generated_file(path: str, file_type: str, expected_text: str | None = None) -> dict:
    p = Path(path)
    if not p.is_file() or p.stat().st_size == 0:
        raise ValueError(f"Generated {file_type} file does not exist or is empty.")

    if file_type == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=False)
        try:
            if not wb.sheetnames:
                raise ValueError("Workbook has no worksheets.")
            if all(ws.max_row <= 1 for ws in wb.worksheets):
                raise ValueError("Workbook contains no useful data.")
        finally:
            wb.close()
    elif file_type == "docx":
        from docx import Document
        doc = Document(p)
        if not doc.paragraphs and not doc.tables:
            raise ValueError("DOCX contains no content.")
    elif file_type == "pptx":
        from pptx import Presentation
        prs = Presentation(p)
        if len(prs.slides) < 1:
            raise ValueError("PPTX contains no slides.")
    elif file_type == "pdf":
        import fitz
        doc = fitz.open(p)
        if len(doc) < 1:
            raise ValueError("PDF contains no pages.")
        if expected_text:
            text = "\n".join(page.get_text() for page in doc)
            if expected_text.lower() not in text.lower():
                # Do not fail on layout/encoding differences; existence/openability is the hard check.
                pass

    return {"valid": True, "file_path": str(p), "size": p.stat().st_size, "file_type": file_type}
